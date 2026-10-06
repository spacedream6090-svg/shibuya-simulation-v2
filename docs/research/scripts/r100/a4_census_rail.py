"""R-100 A4: 経済センサス-活動調査 R3 第30表(町丁目 × 産業中分類・民営)から、
職務の役割に関わる中分類(鉄道業 42・道路旅客運送業 43・道路貨物運送業 44・
運輸に附帯するサービス業 48・郵便業 49・その他の事業サービス業 92 など)を
町丁目ごとに抜き出す。読むだけ。

入力: data/realworld/census_r3/raw/ka21_chu.xlsx(W6/W16 の出所 station_area_industry.json の元表)
実行: システムの python(openpyxl が要る。.venv には無い)
    python docs/research/scripts/r100/a4_census_rail.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
XLSX = ROOT / "data" / "realworld" / "census_r3" / "raw" / "ka21_chu.xlsx"
OUT = ROOT / "data" / "research_cache" / "r100" / "a4_census_rail.json"
AREA13 = ["渋谷１丁目", "渋谷２丁目", "渋谷３丁目", "渋谷４丁目", "道玄坂１丁目", "道玄坂２丁目",
          "宇田川町", "神南１丁目", "神南２丁目", "桜丘町", "南平台町", "円山町", "神泉町"]
CODES = ["42", "43", "44", "47", "48", "49", "86", "92"]


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    codes_row = rows[3]
    names_row = rows[4]
    col = {}
    for j, c in enumerate(codes_row):
        if c is not None and str(c) in CODES:
            col[str(c)] = j
    labels = {k: names_row[j] for k, j in col.items()}
    # 町名・丁目の行を辿る。列 1 = 町名、列 2 = 丁目(無い町は町名の行が末端)。
    out: dict[str, dict] = {}
    town = None
    current = None
    for i, r in enumerate(rows[6:], start=6):
        if r[0] == "渋谷区":
            current = "渋谷区"
        elif r[1]:
            town = str(r[1]).strip()
            current = town
        elif r[2]:
            current = str(r[2]).strip()  # 丁目の行は「渋谷１丁目」の形で町名を含む
        kind = (r[3] or "").strip() if isinstance(r[3], str) else r[3]
        sex = r[4]
        if current is None:
            continue
        rec = out.setdefault(current, {})
        if kind == "事業所数":
            rec["est"] = {k: r[j] for k, j in col.items()} | {"total": r[5]}
        elif kind == "従業者数":
            rec["emp"] = {k: r[j] for k, j in col.items()} | {"total": r[5]}
    res = {"labels": labels, "ward": out.get("渋谷区"), "area13": {}}
    for name in AREA13:
        res["area13"][name] = out.get(name)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print("labels", labels)
    print("ward", out.get("渋谷区"))
    tot = {k: 0 for k in CODES}
    for name in AREA13:
        rec = out.get(name)
        print(name, rec)
        if rec and "emp" in rec:
            for k in CODES:
                v = rec["emp"].get(k)
                if isinstance(v, (int, float)):
                    tot[k] += v
    print("area13 employees by code (x/- を 0 として足した値)", tot)
    print("町丁目の名前(参考・先頭 120)", [k for k in out][:120])


if __name__ == "__main__":
    sys.exit(main())
