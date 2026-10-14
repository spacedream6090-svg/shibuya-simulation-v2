"""R-116 段 0「材料」: 国勢調査の区の表(xlsx)から渋谷区の行を読む(読むだけ)。

入力(すべて構築側=パターン台帳 §10 に無い表・e-Stat):
- census2025_b02_01_*.xlsx    令和7年 人口等基本集計 第2-1表(男女・年齢各歳・国籍総数か日本人)
- census2025_imputed_b02_01_*.xlsx 同 第2-1表の参考表「不詳補完値」(CANCEIS)
- census2025_b10_*.xlsx       令和7年 第10表(世帯の家族類型 × 世帯人員の人数)
- census2025_b13_01_01_*.xlsx 令和7年 第13-1表(男女・年齢5歳・続き柄・配偶関係・家族類型)
- census2020_e11_01_*.xlsx    令和2年 従業地・通学地集計 第11-1表(従業地の男女・年齢5歳・産業大分類)
出力: docs/research/scripts/r116/r116_ward_xlsx.json

注: .venv に openpyxl が無いので、openpyxl の入った python(3.12)で動かす:
    py -3.12 docs/research/scripts/r116/r116_ward_xlsx.py
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
CACHE = ROOT / "data" / "research_cache" / "r116"
OUT = Path(__file__).with_suffix(".json")
NAME = "13113_渋谷区"


def md5(p: Path) -> str:
    h = hashlib.md5()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def n(v):
    if v in (None, "-", "…", "***", "X"):
        return 0
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def read(p: Path, head_idx: int, first_val: int, name_cols: range):
    """head_idx 行目の項目名(first_val 列から)と、渋谷区の行を返す。"""
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb[wb.sheetnames[0]]
    labels, rows = None, []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == head_idx:
            labels = list(row[first_val:])
        if i > head_idx and any(j < len(row) and row[j] == NAME for j in name_cols):
            rows.append(list(row))
    wb.close()
    return labels, rows


def age_of(lab: str):
    m = re.match(r"\d+_(\d+)歳", lab)
    return int(m.group(1)) if m else None


def b02(p: Path) -> dict:
    labels, rows = read(p, 7, 4, range(0, 4))
    res = {}
    for r in rows:
        vals = dict(zip(labels, r[4:]))
        agg = {"total": n(vals.get("000_総数")), "age0_14": 0, "age15_64": 0, "age65p": 0}
        for lab, v in vals.items():
            if not isinstance(lab, str):
                continue
            a = age_of(lab) if lab[:1].isdigit() else None  # 再掲(R…)は足さない
            if a is None:
                continue
            k = "age0_14" if a <= 14 else "age15_64" if a <= 64 else "age65p"
            agg[k] += n(v) or 0
        agg["age_unknown"] = sum(n(v) or 0 for lab, v in vals.items() if isinstance(lab, str) and "年齢「不詳」" in lab)
        res[f"{r[0]}|{r[1]}"] = agg
    return {"shibuya": res}


def b10(p: Path) -> dict:
    labels, rows = read(p, 6, 4, range(0, 4))
    return {"shibuya_by_size": {str(r[3]): {lab: n(v) for lab, v in zip(labels, r[4:]) if isinstance(lab, str)}
                                for r in rows}}


def b13(p: Path) -> dict:
    labels, rows = read(p, 6, 6, range(0, 3))
    tot = None
    for r in rows:
        if r[2] == "0_総数" and r[3] == "0_総数" and r[4] == "0_総数" and r[5] == "00_総数":
            tot = {lab: n(v) for lab, v in zip(labels, r[6:]) if isinstance(lab, str)}
    dims = {"配偶関係": sorted({str(r[2]) for r in rows}), "男女": sorted({str(r[3]) for r in rows}),
            "世帯の家族類型": sorted({str(r[4]) for r in rows}), "年齢": sorted({str(r[5]) for r in rows})}
    return {"n_rows_shibuya": len(rows), "relation_columns": [x for x in labels if isinstance(x, str)],
            "dims": {k: [len(v), v] for k, v in dims.items()}, "shibuya_total_by_relation": tot}


def e11(p: Path) -> dict:
    labels, rows = read(p, 7, 5, range(0, 4))
    d = {f"{r[3]}|{r[4]}": {lab: n(v) for lab, v in zip(labels, r[5:]) if isinstance(lab, str)} for r in rows}
    tot = d.get("0_総数|00_総数", {})
    return {"workplace_total": tot.get("0_総数"),
            "by_sex": {k.split("|")[0]: v.get("0_総数") for k, v in d.items() if k.endswith("|00_総数")},
            "by_age": {k.split("|")[1]: v.get("0_総数") for k, v in d.items() if k.startswith("0_総数|")},
            "by_industry": {k: v for k, v in tot.items() if k != "0_総数"}}


def main() -> int:
    files = {
        "b02_01": next(CACHE.glob("census2025_b02_01_*.xlsx")),
        "b02_01_imputed": next(CACHE.glob("census2025_imputed_b02_01_*.xlsx")),
        "b10": next(CACHE.glob("census2025_b10_*.xlsx")),
        "b13_01_01": next(CACHE.glob("census2025_b13_01_01_*.xlsx")),
        "e11_01": next(CACHE.glob("census2020_e11_01_*.xlsx")),
    }
    out = {"md5": {k: [p.name, md5(p)] for k, p in files.items()},
           "census2025_b02_01": b02(files["b02_01"]),
           "census2025_imputed_b02_01": b02(files["b02_01_imputed"]),
           "census2025_b10": b10(files["b10"]),
           "census2025_b13_01_01": b13(files["b13_01_01"]),
           "census2020_e11_01": e11(files["e11_01"])}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8", newline="\n")
    print(json.dumps(out, ensure_ascii=False, indent=0)[:7000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
