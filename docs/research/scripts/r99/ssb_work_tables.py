"""R-99 A1/A2: summarise SSB 2021 tables (read only).

Reads data/research_cache/r99/*.xlsx and prints:
 (1) table 74-9: Tokyo, 'work' participation rate by occupation x day (weekday/Sat/Sun)
 (2) table 14: national, weekday, 'work' by occupation x 15-min slot -> start proxy quantiles
     (positive increments in 04:00-12:00, same method as v2-d68-remaining-research.md section 1-1),
     peak rate, rate at 0:00 and 22:00
 (3) tables 15-1/2/3: Tokyo, employment status x day, 'work' -> peak rate and start proxy
Usage: python docs/research/scripts/r99/ssb_work_tables.py
"""
import os, sys
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
D = os.path.join(os.path.dirname(__file__), "..", "..", "..", "..", "data", "research_cache", "r99")

def num(x):
    if x in (None, "-", "…", "x", "X", "***"): return 0.0
    try: return float(x)
    except Exception: return 0.0

def slot_label(i):
    m = i * 15
    return f"{m//60}:{m%60:02d}"

def start_proxy(v):
    # v: 96 slot rates. increments over slots whose start is in [4:00, 12:00)
    inc = []
    for i in range(16, 48):
        inc.append((i, max(0.0, v[i] - v[i-1])))
    tot = sum(d for _, d in inc)
    if tot <= 0: return tot, None, None, None, None
    cum = 0; q = {}
    for i, d in inc:
        cum += d
        for p in (0.1, 0.5, 0.9):
            if p not in q and cum >= p * tot: q[p] = slot_label(i)
    r = sum(d for i, d in inc if (i * 15) % 30 == 0) / tot
    return tot, q[0.1], q[0.5], q[0.9], r

def t74_9():
    wb = openpyxl.load_workbook(os.path.join(D, "ssb2021_seikatsu_t74-9_occupation_pref.xlsx"), read_only=True)
    ws = wb.worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=10, values_only=True):
        day, area, sex, occ = row[0], row[1], row[2], row[3]
        if area is None or not str(area).startswith(("00_", "13_")) or sex != "0_総数": continue
        out[(area, occ, day)] = (num(row[4]), num(row[9]))  # pop(thousand), work rate
    print("== (1) 74-9 work participation rate % (pop thousand) [weekday / Sat / Sun]")
    occs = []
    for k in out:
        if k[1] not in occs: occs.append(k[1])
    for area in ("00_全国", "13_東京都"):
        print(" area", area)
        for occ in occs:
            vals = []
            for day in ("2_平日", "3_土曜日", "4_日曜日"):
                p, w = out.get((area, occ, day), (0, 0))
                vals.append(f"{w:5.1f}")
            p = out.get((area, occ, "1_週全体"), (0, 0))[0]
            print(f"   {occ:<40} pop={p:>7.0f}  " + " / ".join(vals))

def t14():
    wb = openpyxl.load_workbook(os.path.join(D, "ssb2021_jikantai_t14_occupation.xlsx"), read_only=True)
    ws = wb.worksheets[0]
    print("== (2) table 14 national weekday 'work' by occupation: start proxy p10/p50/p90, :00+:30 share, peak, 0:00, 22:00")
    for row in ws.iter_rows(min_row=10, values_only=True):
        if row[0] != "1_平日" or row[2] != "0_総数" or row[4] != "05_仕事": continue
        v = [num(x) for x in row[6:102]]
        tot, a, b, c, r = start_proxy(v)
        pk = max(range(96), key=lambda i: v[i])
        print(f"   {row[3]:<34} inc={tot:6.2f} start p10/50/90={a}/{b}/{c} r00_30={r:.3f} peak={v[pk]:.2f}@{slot_label(pk)} 0:00={v[0]:.2f} 22:00={v[88]:.2f}")

def t15(fn, label):
    wb = openpyxl.load_workbook(os.path.join(D, fn), read_only=True)
    ws = wb.worksheets[0]
    hdr = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True)):
        if row and row[0] == "曜日": hdr = row; break
    print(f"== (3) {label}: header {[h for h in hdr if h]}")
    for row in ws.iter_rows(min_row=10, values_only=True):
        if row[0] is None: continue
        rec = row
        # find columns: day, area, sex, status, activity, pop, 96 slots
        if not str(rec[1]).startswith(("00_", "13_")) or rec[2] != "0_総数" or rec[4] != "05_仕事": continue
        v = [num(x) for x in rec[6:102]]
        tot, a, b, c, r = start_proxy(v)
        pk = max(range(96), key=lambda i: v[i])
        print(f"   {rec[1]:<8} {rec[3]:<40} inc={tot:6.2f} start p10/50/90={a}/{b}/{c} r={r if r is None else round(r,3)} peak={v[pk]:.2f}@{slot_label(pk)} 0:00={v[0]:.2f} 22:00={v[88]:.2f}")

if __name__ == "__main__":
    t74_9()
    t14()
    t15("ssb2021_jikantai_t15-1_emp_weekday.xlsx", "15-1 weekday")
    t15("ssb2021_jikantai_t15-2_emp_sat.xlsx", "15-2 Saturday")
    t15("ssb2021_jikantai_t15-3_emp_sun.xlsx", "15-3 Sunday")
