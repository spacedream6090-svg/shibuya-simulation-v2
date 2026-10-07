"""R-99 A2: annual working days class shares from Employment Status Survey 2022, pref. table 19 (read only).

Reads data/research_cache/r99/shugyo2022_pref_t19_industry_days.xlsx, prints for 東京都 and 特別区部
(if present) the share of 年間就業日数 top classes by employment status and by industry.
Usage: python docs/research/scripts/r99/shugyo_days.py
"""
import os, sys, re
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
F = os.path.join(os.path.dirname(__file__), "..", "..", "..", "..", "data", "research_cache", "r99",
                 "shugyo2022_pref_t19_industry_days.xlsx")

def num(x):
    try: return float(x)
    except Exception: return 0.0

wb = openpyxl.load_workbook(F, read_only=True)
ws = wb.worksheets[0]
rows = ws.iter_rows(values_only=True)
hdr = None
for i, r in enumerate(rows):
    if i == 6: hdr = r
    if i == 8: break
# header names are aligned with the full row (label columns are None in the header)
names = list(hdr)
# top-level classes: codes '0','1','2','3','4' ... single digit before '_'
top = [(j, n) for j, n in enumerate(names) if n and re.match(r"^\d_", str(n))]
print("top classes:", [n for _, n in top])
areas = set()
data = {}
for r in ws.iter_rows(min_row=10, values_only=True):
    area, sex, emp, ind = r[1], r[3], r[5], r[7]
    areas.add(area)
    if sex != "0_総数": continue
    if not (str(area).startswith("13") ): continue
    data[(area, emp, ind)] = [num(r[j]) for j, _ in top]
print("tokyo areas:", sorted(a for a in areas if str(a).startswith("13")))
def show(area, emp, ind):
    v = data.get((area, emp, ind))
    if not v or v[0] == 0: return None
    tot = v[0]
    return f"N={tot:>10.0f} " + " ".join(f"{n.split('_',1)[1]}={x/tot*100:5.1f}%" for (j, n), x in zip(top[1:], v[1:]))
for area in sorted(a for a in areas if str(a).startswith("13")):
    print("====", area)
    emps = []
    for (a, e, i) in data:
        if a == area and e not in emps: emps.append(e)
    for e in emps:
        s = show(area, e, "0_総数")
        if s: print("  emp", f"{e:<36}", s)
    inds = []
    for (a, e, i) in data:
        if a == area and e == "0_総数" and i not in inds: inds.append(i)
    for i in inds:
        s = show(area, "0_総数", i)
        if s: print("  ind", f"{i:<36}", s)
