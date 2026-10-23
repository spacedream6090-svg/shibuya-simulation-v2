"""第 4 期 4-1 段 1〜3(W21)の材料の写し: 国勢調査の xlsx から渋谷区の行だけを JSON に写す(読むだけ)。

.venv に openpyxl が無いので、openpyxl の入った python 3.12(システム)で動かす(発注の規律):

    py -3.12 docs/bench/analysis/wallbounce-1012/pop41/tools/xlsx_to_derived.py

- 入力: ``data/research_cache/r116/`` と ``data/research_cache/pop41/`` の xlsx(どれも構築側=パターン台帳 §10 の
  検証専用の一覧に無い表・R-116 §1 の「構築」の行)。下の ``SPECS`` に無いファイルは開かない。
- 出力: ``data/research_cache/pop41/derived/<名>.json``(gitignore の下)。形は
  ``{schema, source, source_md5, sheet, title, dims, items, rows}``・``rows`` は ``[次元の値..., [値...]]``。
  値の写し方: 数はそのまま・``-``(該当なし)は 0・``X``/``…``/``***``(秘匿・不詳)は null。
- 決定的: 同じ xlsx から同じバイト列(キー整列・区切りの空白なし・UTF-8・末尾に改行 1 つ・時刻を入れない)。
- 表の形の見つけ方: 「地域識別コード」で始まる行=次元の名の行。その行の空白 1 字のセルの列から右が値の列。
  値の列の項目名は、その行より上で値の列の頭のセルが ``<番号>_<名>`` の形の最も近い行。
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[6]
R116 = ROOT / "data" / "research_cache" / "r116"
POP41 = ROOT / "data" / "research_cache" / "pop41"
OUT = POP41 / "derived"
WARD = "13113_渋谷区"

#: (入力の置き場, ファイル名, 写しの名)
SPECS: tuple[tuple[Path, str, str], ...] = (
    (R116, "census2025_b02_01_age_single_sex_wards.xlsx", "c25_b02_01"),
    (R116, "census2025_imputed_b02_01_age_single_sex_wards.xlsx", "c25w_b02_01"),
    (R116, "census2025_b10_family_type_hhsize_muni.xlsx", "c25_b10"),
    (R116, "census2025_b13_01_01_sex_age_relation_marital_family_wards.xlsx", "c25_b13_01"),
    (R116, "census2020_e11_01_workplace_sex_age_industry_muni.xlsx", "c20_e11_01"),
    (POP41, "census2025_b12_01_head_sex_age_marital_family_wards.xlsx", "c25_b12_01"),
    (POP41, "census2025_b15_01_husband_age_children_wards.xlsx", "c25_b15_01"),
    (POP41, "census2025_b19_01_husband_wife_age_wards.xlsx", "c25_b19_01"),
    (POP41, "census2025_imputed_b10_family_type_hhsize_muni.xlsx", "c25w_b10"),
    (POP41, "census2025_imputed_b13_01_01_sex_age_relation_marital_family_wards.xlsx", "c25w_b13_01"),
    (POP41, "census2025_imputed_b13_01_02_male.xlsx", "c25w_b13_02"),
    (POP41, "census2025_imputed_b13_01_03_female.xlsx", "c25w_b13_03"),
    (POP41, "census2020r_imputed_b02_01_age_single_sex_wards.xlsx", "c20w_b02_01"),
    (POP41, "census2020r_imputed_b10_family_type_hhsize_muni.xlsx", "c20w_b10"),
    (POP41, "census2020r_imputed_b13_01_01_sex_age_relation_marital_family_wards.xlsx", "c20w_b13_01"),
    (POP41, "census2020r_imputed_b13_01_02_male.xlsx", "c20w_b13_02"),
    (POP41, "census2020r_imputed_b13_01_03_female.xlsx", "c20w_b13_03"),
    (POP41, "census2020_c03_01_age_marital_labour_status_wards.xlsx", "c20_c03_01"),
    (POP41, "census2020_c05_01_age_status_industry_wards.xlsx", "c20_c05_01"),
    (POP41, "census2020_c10_01_age_industry_occupation_wards.xlsx", "c20_c10_01"),
    (POP41, "census2020_e08_01_status_industry_workplace_muni.xlsx", "c20_e08_01"),
    (POP41, "census2020_e15_01_age_occupation_workplace_muni.xlsx", "c20_e15_01"),
    (POP41, "census2020_e16_01_industry_occupation_workplace_muni.xlsx", "c20_e16_01"),
    (POP41, "census2020_e01_03_students_sex_age_place_muni.xlsx", "c20_e01_03"),  # 〔検収 L6〕外から来る通学者
)
_LABEL = re.compile(r"^[0-9A-Za-z]+_")


def md5(p: Path) -> str:
    h = hashlib.md5()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def cell(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else v
    s = str(v).strip()
    if s in ("-", ""):
        return 0
    if s in ("X", "x", "…", "***", "..."):
        return None
    try:
        return int(s.replace(",", ""))
    except ValueError:
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return s


def convert(src: Path) -> dict:
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    title = next((str(r[0]) for r in rows[:4] if r and isinstance(r[0], str) and "表" in r[0][:6]), "")
    if not title:
        title = next((str(c) for r in rows[:4] for c in r if isinstance(c, str) and "表" in c[:6]), "")
    di = next(i for i, r in enumerate(rows) if any(isinstance(c, str) and c.startswith("地域識別コード") for c in r[:6]))
    drow = rows[di]
    fv = next(j for j, c in enumerate(drow) if isinstance(c, str) and c == " ")
    dims = [str(c) for c in drow[:fv]]
    ii = next(i for i in range(di - 1, -1, -1) if fv < len(rows[i]) and isinstance(rows[i][fv], str)
              and _LABEL.match(rows[i][fv]))
    items = []
    for c in rows[ii][fv:]:
        if not isinstance(c, str):
            break
        items.append(c)
    area_cols = [j for j, d in enumerate(dims) if d.startswith(("地域名", "地域識別"))]
    out_rows = []
    for r in rows[di + 1:]:
        if not r or not any(j < len(r) and r[j] == WARD for j in area_cols):
            continue
        dv = [str(r[j]) if r[j] is not None else "" for j in range(fv)]
        vals = [cell(r[fv + k]) if fv + k < len(r) else None for k in range(len(items))]
        out_rows.append(dv + [vals])
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": ws.title, "title": title, "dims": dims, "items": items, "rows": out_rows}


def main(argv: list[str] | None = None) -> int:
    only = set(sys.argv[1:] if argv is None else argv)  # 写しの名を並べればその表だけ(既定=全部)
    OUT.mkdir(parents=True, exist_ok=True)
    for base, name, stem in SPECS:
        if only and stem not in only:
            continue
        src = base / name
        doc = convert(src)
        data = json.dumps(doc, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8") + b"\n"
        dst = OUT / f"{stem}.json"
        dst.write_bytes(data)
        print(f"{stem}: rows {len(doc['rows'])} items {len(doc['items'])} dims {len(doc['dims'])} "
              f"md5 {hashlib.md5(data).hexdigest()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
