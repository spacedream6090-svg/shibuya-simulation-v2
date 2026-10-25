"""第 4 期 4-1 段 4〜7(W21 の続き)の材料の写し: xlsx から使う行だけを JSON に写す(読むだけ)。

.venv に openpyxl が無いので、openpyxl の入った python 3.12(システム)で動かす(発注の規律):

    py -3.12 docs/bench/analysis/wallbounce-1012/pop41b/tools/xlsx_to_derived_b.py [写しの名 ...]

- 入力: 下の ``SPECS`` のファイルだけ(どれも構築側=パターン台帳 §10 の検証専用の一覧に無い表・R-116 §1-4・§1-5 の
  「構築」の行)。``SPECS`` に無いファイルは開かない。
- 出力: ``data/research_cache/pop41b/derived/<名>.json``(gitignore の下)。形は段 1〜3 の写しと同じ
  ``{schema, source, source_md5, sheet, title, dims, items, rows}``・``rows`` は ``[次元の値..., [値...]]``。
  値の写し方: 数はそのまま・``-``(該当なし)は 0・``X``/``x``/``…``/``***``(秘匿・不詳)は null。
- 決定的: 同じ xlsx から同じバイト列(キー整列・区切りの空白なし・UTF-8・末尾に改行 1 つ・時刻を入れない)。
- 形は 3 種類: ``census``(「地域識別コード」の見出しの行がある e-Stat の表・地域の列で行を絞る)・``ecen_t28``
  (経済センサス 第28表=町丁目 × 産業大分類)・``wage``(賃金構造基本統計調査の 4 表=企業規模計の欄だけ)。
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Any, Callable

import openpyxl

ROOT = Path(__file__).resolve().parents[6]
DATA = ROOT / "data"
OUT = DATA / "research_cache" / "pop41b" / "derived"
WARD = "13113_渋谷区"
_LABEL = re.compile(r"^[0-9A-Za-z]+_")


def md5(p: Path) -> str:
    h = hashlib.md5()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else v
    s = str(v).strip()
    if s in ("-", ""):
        return 0
    if s in ("X", "x", "…", "***", "..."):
        return None
    try:
        return int(s.replace(",", ""))
    except ValueError:
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return s


def _rows(src: Path, sheet: int = 0) -> tuple[str, list[list[Any]]]:
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[wb.sheetnames[sheet]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    title = ws.title
    wb.close()
    return title, rows


def _title(rows: list[list[Any]]) -> str:
    for r in rows[:6]:
        for c in r:
            if isinstance(c, str) and re.search(r"第[０-９0-9－\-]+表", c[:24]) and len(c) > 6:
                return c
    return ""


# ---------------------------------------------------------------------------------------------- census
def census(src: Path, keep: Callable[[dict[str, str]], bool], item_rows: int = 1) -> dict[str, Any]:
    """「地域識別コード」の行=次元の名。その行の空白 1 字のセルから右が値の列。``item_rows`` 段の項目名を「|」で結ぶ。"""
    sheet, rows = _rows(src)
    di = next(i for i, r in enumerate(rows) if any(isinstance(c, str) and c.startswith("地域識別コード") for c in r[:10]))
    drow = rows[di]
    fv = next(j for j, c in enumerate(drow) if isinstance(c, str) and c == " ")
    dims = [str(c) for c in drow[:fv]]
    ii = next(i for i in range(di - 1, -1, -1) if fv < len(rows[i]) and isinstance(rows[i][fv], str)
              and _LABEL.match(rows[i][fv]))
    n_items = 0
    for c in rows[ii][fv:]:
        if not isinstance(c, str):
            break
        n_items += 1
    parts = [rows[ii + k][fv:fv + n_items] for k in range(item_rows)]
    items = ["|".join(str(p[j]) for p in parts) for j in range(n_items)]
    out_rows = []
    for r in rows[di + 1:]:
        if not r or r[0] is None and all(c is None for c in r[:fv]):
            continue
        dv = [str(r[j]) if j < len(r) and r[j] is not None else "" for j in range(fv)]
        if not keep(dict(zip(dims, dv))):
            continue
        vals = [cell(r[fv + k]) if fv + k < len(r) else None for k in range(n_items)]
        out_rows.append(dv + [vals])
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": sheet, "title": _title(rows), "dims": dims, "items": items, "rows": out_rows}


# ---------------------------------------------------------------------------------------------- 第28表
def ecen_t28(src: Path) -> dict[str, Any]:
    """経済センサス 2021 第28表(区 → 町 → 丁目の行・産業大分類ごとに事業所数と従業者数)。``level`` は ward/town/chome。"""
    sheet, rows = _rows(src)
    head = rows[3]
    majors = []
    j = 5
    while j < len(head):
        lab = head[j]
        if isinstance(lab, str) and lab.strip():
            letter = lab.strip().split()[0].split("\n")[0].strip()
            majors.append((letter[0], j))
        j += 2
    items = ["総数_事業所数", "総数_従業者数"]
    for m, _j in majors:
        items += [f"{m}_事業所数", f"{m}_従業者数"]
    out_rows = []
    town = ""
    for r in rows[6:]:
        if not r or all(c is None for c in r[:5]):
            continue
        if r[0]:
            level, name = "ward", str(r[0]).strip()
            town = ""
        elif r[1]:
            level, name = "town", str(r[1]).strip()
            town = name
        elif r[2]:
            level, name = "chome", str(r[2]).strip()
        else:
            continue
        vals = [cell(r[3]), cell(r[4])]
        for _m, jj in majors:
            vals += [cell(r[jj]) if jj < len(r) else None, cell(r[jj + 1]) if jj + 1 < len(r) else None]
        if all(v is None for v in vals):  # 注の行
            continue
        out_rows.append([level, town if level == "chome" else name, name if level == "chome" else "", vals])
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": sheet, "title": _title(rows), "dims": ["level", "town", "chome"], "items": items,
            "rows": out_rows}


# ---------------------------------------------------------------------------------------------- 賃金構造
_W8 = ["年齢", "勤続年数", "所定内実労働時間数", "超過実労働時間数", "きまって支給する現金給与額", "所定内給与額",
       "年間賞与その他特別給与額", "労働者数"]


def _clean(s: Any) -> str:
    return re.sub(r"[\s　]+", "", str(s or ""))


def wage_t9(src: Path) -> dict[str, Any]:
    """第9表(職種 × 性 × 雇用形態)・企業規模計の 8 欄。区画の頭の行は「職種\\n性\\n常用労働者計」か「性\\n常用労働者計」。"""
    sheet, rows = _rows(src)
    occ = sex = ""
    out_rows = []
    for r in rows[10:]:
        lab = r[1] if len(r) > 1 else None
        if not isinstance(lab, str) or not lab.strip():
            continue
        parts = [_clean(p) for p in lab.split("\n") if _clean(p)]
        if len(parts) == 3:
            occ, sex, cat = parts
        elif len(parts) == 2:
            sex, cat = parts
        else:
            cat = parts[0]
        out_rows.append([occ, sex, cat, [cell(c) for c in r[3:11]]])
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": sheet, "title": _title(rows), "dims": ["職種", "性", "雇用形態"], "items": _W8, "rows": out_rows}


def wage_t6(src: Path) -> dict[str, Any]:
    """第6表(職種 × 性 × 年齢)・3 枚(男女計・男・女)・企業規模計の 8 欄。"""
    wb = openpyxl.load_workbook(src, read_only=True)
    out_rows = []
    title = ""
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        title = title or _title(rows)
        occ = ""
        for r in rows[11:]:
            lab = r[1] if len(r) > 1 else None
            if not isinstance(lab, str) or not lab.strip():
                continue
            c = _clean(lab)
            if "歳" in c:
                out_rows.append([sn, occ, c, [cell(x) for x in r[3:11]]])
            else:
                occ = c
                out_rows.append([sn, occ, "計", [cell(x) for x in r[3:11]]])
    wb.close()
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": "+".join(["男女計", "男", "女"]), "title": title, "dims": ["性", "職種", "年齢"], "items": _W8,
            "rows": out_rows}


def wage_t18(src: Path) -> dict[str, Any]:
    """第18表(職種 × 性 × 所定内給与額の階級と分布特性値)・3 枚・企業規模計の区画だけ。項目=職種(列)。"""
    wb = openpyxl.load_workbook(src, read_only=True)
    out_rows = []
    items: list[str] = []
    title = ""
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        title = title or _title(rows)
        hi = next(i for i, r in enumerate(rows) if len(r) > 1 and _clean(r[1]) == "区分")
        occs = []
        for j in range(3, len(rows[hi])):
            v = rows[hi][j]
            if not isinstance(v, str) or not v.strip():
                break
            occs.append((j, _clean(v)))
        if not items:
            items = [o for _j, o in occs]
        started = False
        for r in rows[hi + 1:]:
            lab = _clean(r[1] if len(r) > 1 else "")
            if not lab:
                continue
            if "企業規模計" in lab:
                started = True
                out_rows.append([sn, "計", [cell(r[j]) for j, _o in occs]])
                continue
            if not started:
                continue
            if "人以上" in lab or "人～" in lab or "～999人" in lab or "～99人" in lab:
                break
            out_rows.append([sn, lab, [cell(r[j]) for j, _o in occs]])
    wb.close()
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": "男女計+男+女", "title": title, "dims": ["性", "区分"], "items": items, "rows": out_rows}


_T2 = ["年齢", "勤続年数", "実労働日数", "1日当たり所定内実労働時間数", "1時間当たり所定内給与額", "年間賞与その他特別給与額",
       "労働者数"]


def wage_t2(src: Path) -> dict[str, Any]:
    """短時間労働者 第2表(職種 × 性)・企業規模計の区画だけ。項目=男女計・男・女 × 7 欄。"""
    sheet, rows = _rows(src)
    items = [f"{s}|{k}" for s in ("男女計", "男", "女") for k in _T2]
    out_rows = []
    started = False
    for r in rows[9:]:
        lab = r[1] if len(r) > 1 else None
        if not isinstance(lab, str) or not lab.strip():
            continue
        parts = [_clean(p) for p in lab.split("\n") if _clean(p)]
        if len(parts) >= 2:
            if started:
                break
            started = "企業規模計" in parts[0]
            occ = parts[-1]
        else:
            occ = parts[0]
        if started:
            out_rows.append([occ, [cell(c) for c in r[3:24]]])
    return {"schema": "w21.derived/1", "source": src.relative_to(ROOT).as_posix(), "source_md5": md5(src),
            "sheet": sheet, "title": _title(rows), "dims": ["職種"], "items": items, "rows": out_rows}


# ---------------------------------------------------------------------------------------------- 表の一覧
R107 = DATA / "research_cache" / "r107"
R109 = DATA / "research_cache" / "r109"
POP41 = DATA / "research_cache" / "pop41"
POP41B = DATA / "research_cache" / "pop41b"
CENSUS_R3 = DATA / "realworld" / "census_r3" / "raw"


def _not_old(d: dict[str, str]) -> bool:
    return "旧" not in d.get("地域名", "")


SPECS: dict[str, Callable[[], dict[str, Any]]] = {
    # 経済センサス 2021 第28表(町丁目 × 産業大分類・民営)=舞台の席の配分と従業者の割合
    "ecen21_t28": lambda: ecen_t28(CENSUS_R3 / "ka21_dai.xlsx"),
    # 同 第6-2表(区 × 産業中分類 × 従業上の地位 6 区分 × 男女)・渋谷区の行だけ
    "ecen21_t6_2": lambda: census(POP41B / "ecensus2021_t6-2_industry_status_muni.xlsx",
                                  lambda d: d.get("地域区分") == WARD, item_rows=2),
    # 国勢調査 2020 従業地・通学地 第4表(東京都=従業地)・従業地が渋谷区の行だけ(常住地は全国の市区町村)
    "c20_t04_in": lambda: census(R109 / "census2020_t04_13_workplace_by_residence.xlsx",
                                 lambda d: d.get("地域名（従業地・通学地）") == WARD),
    # 国勢調査 2025 第2-5表【不詳補完値】(市区町村 × 男女 × 各歳)・国籍総数・男女別・旧市町村を除く
    "c25w_b02_05": lambda: census(POP41B / "census2025_imputed_b02_05_age_single_sex_muni.xlsx",
                                  lambda d: d.get("国籍総数か日本人") == "0_国籍総数" and d.get("男女") in ("1_男", "2_女")
                                  and _not_old(d)),
    # 国勢調査 2025 第12-3表【不詳補完値】(市区町村 × 世帯主の年齢 × 家族類型)・世帯主の男女=総数
    "c25w_b12_03": lambda: census(POP41B / "census2025_imputed_b12_03_head_family_muni.xlsx",
                                  lambda d: d.get("世帯主の男女") == "0_総数"),
    # 国勢調査 2025 第10表【不詳補完値】(市区町村 × 家族類型 × 世帯人員)・全部の地域
    "c25w_b10_muni": lambda: census(POP41 / "census2025_imputed_b10_family_type_hhsize_muni.xlsx", lambda d: True),
    # 賃金構造基本統計調査 令和 7 年(R-107)
    "wage_t9": lambda: wage_t9(R107 / "chinkou_r7_ippan_shokushu_t9_occ_sex_emptype.xlsx"),
    "wage_t6": lambda: wage_t6(R107 / "chinkou_r7_ippan_shokushu_t6_occ_sex_age.xlsx"),
    "wage_t18": lambda: wage_t18(R107 / "chinkou_r7_ippan_shokushu_t18_occ_sex_wagedist.xlsx"),
    "wage_t2": lambda: wage_t2(R107 / "chinkou_r7_tanjikan_shokushu_t2_hourly.xlsx"),
}


def main(argv: list[str] | None = None) -> int:
    only = set(sys.argv[1:] if argv is None else argv)
    OUT.mkdir(parents=True, exist_ok=True)
    for stem, fn in SPECS.items():
        if only and stem not in only:
            continue
        doc = fn()
        data = json.dumps(doc, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8") + b"\n"
        (OUT / f"{stem}.json").write_bytes(data)
        print(f"{stem}: rows {len(doc['rows'])} items {len(doc['items'])} dims {len(doc['dims'])} "
              f"src_md5 {doc['source_md5']} md5 {hashlib.md5(data).hexdigest()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
